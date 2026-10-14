#!/usr/bin/env python3
"""
레시피 Excel에서 메뉴별 알러지 정보를 추출합니다.
출력: data/menu_allergens.json, data/allergen_report.xlsx
"""

import json
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from allergens import ALLERGEN_NAMES, KOREAN_ALLERGENS

EXCEL_PATH = Path("documents/[CON] 2026 셀렉션  가을&겨울 레시피 (1).xlsx")
SKIP_SHEETS = {"레시피 작성 가이드", "시트5"}

# 좌측 블록: B(1)~H(7), 우측 블록: I(8)~O(14)
L = {"name": 1, "batch_marker": 3, "mult": 4, "unit": 5, "storage": 7}
R = {"name": 8, "batch_marker": 10, "mult": 11, "unit": 12, "storage": 14}


# ─────────────────────────────────────────────
# 1. Excel 파싱
# ─────────────────────────────────────────────

def _extract_ingredients_from_block(rows, start, side):
    """Components 헤더 이후부터 로스량 행까지 재료명 수집."""
    n = L["name"] if side == "L" else R["name"]
    u = L["unit"] if side == "L" else R["unit"]
    end_marker = "로스량(g)"

    ingredients = []
    comp_found = False
    for row in rows[start:]:
        if row[n] == "Components":
            comp_found = True
            continue
        if not comp_found:
            continue
        if row[n] == end_marker:
            break
        if row[n] and row[u]:
            ingredients.append(str(row[n]).strip())
    return ingredients


def parse_sheet(ws) -> dict[str, list[str]]:
    """시트 → {블록명: [원재료명, ...]} 딕셔너리 반환."""
    rows = list(ws.iter_rows(values_only=True))
    blocks: dict[str, list[str]] = {}

    for i, row in enumerate(rows):
        left_name = str(row[L["name"]]).strip() if (row[L["name"]] and row[L["batch_marker"]] == "배합수") else None
        right_name = str(row[R["name"]]).strip() if (row[R["name"]] and row[R["batch_marker"]] == "배합수") else None

        if left_name:
            blocks[left_name] = _extract_ingredients_from_block(rows, i + 1, "L")
        if right_name:
            blocks[right_name] = _extract_ingredients_from_block(rows, i + 1, "R")

    return blocks


# ─────────────────────────────────────────────
# 2. 재귀 해소 (서브레시피 → 원재료)
# ─────────────────────────────────────────────

_PREFIXES = ["(프랩)", "(플레이팅)", "프랩_", "프랩 "]


def _normalize(name: str) -> str:
    n = name.strip().lower()
    for p in _PREFIXES:
        if n.startswith(p.lower()):
            n = n[len(p):].strip()
    return n.replace(" ", "")


def _find_block(ing_name: str, blocks: dict) -> str | None:
    """(프랩) 접두어·공백 차이를 무시하고 블록명 매칭."""
    if ing_name in blocks:
        return ing_name
    target = _normalize(ing_name)
    for block_name in blocks:
        if _normalize(block_name) == target:
            return block_name
    return None


def collect_raw_ingredients(block_name: str, blocks: dict, visited: set | None = None) -> list[str]:
    if visited is None:
        visited = set()
    if block_name in visited:
        return []
    visited.add(block_name)

    actual = _find_block(block_name, blocks)
    if actual is None:
        return [block_name]

    result = []
    for ing in blocks[actual]:
        matched = _find_block(ing, blocks)
        if matched:
            result.extend(collect_raw_ingredients(matched, blocks, set(visited)))
        else:
            result.append(ing)
    return result


def get_top_block(blocks: dict) -> str | None:
    """완성계량 또는 플레이팅 블록 반환. 없으면 None → 전체 수집 폴백."""
    for name in blocks:
        if "완성계량" in name:
            return name
    for name in blocks:
        if "플레이팅" in name:
            return name
    return None


# ─────────────────────────────────────────────
# 3. 알러지 매핑
# ─────────────────────────────────────────────

def map_allergens(ingredients: list[str]) -> dict[str, list[str]]:
    """원재료 → {알러지명: [유발 재료, ...]} 매핑."""
    found: dict[str, list[str]] = {}
    for ing in ingredients:
        ing_lower = ing.lower()
        for allergen, keywords in KOREAN_ALLERGENS.items():
            for kw in keywords:
                if kw in ing_lower:
                    found.setdefault(allergen, []).append(ing)
                    break
    return found


def get_category(sheet_name: str) -> str:
    if sheet_name == "피클":
        return "피클/기타"
    if "잠실" in sheet_name:
        return "잠실 단품"
    if "단품" in sheet_name:
        return "단품"
    if "코스" in sheet_name:
        return "코스"
    return "기타"


# ─────────────────────────────────────────────
# 4. JSON 생성
# ─────────────────────────────────────────────

def collect_all_ingredients(blocks: dict) -> list[str]:
    """완성계량 없는 시트: 모든 블록 원재료를 합산."""
    all_raw = []
    for name in blocks:
        all_raw.extend(collect_raw_ingredients(name, blocks))
    return list(set(all_raw))


def build_menu_allergens(wb) -> dict:
    result = {}
    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        blocks = parse_sheet(ws)
        if not blocks:
            continue

        top = get_top_block(blocks)
        if top:
            raw = collect_raw_ingredients(top, blocks)
        else:
            raw = collect_all_ingredients(blocks)

        allergen_map = map_allergens(raw)

        result[sheet_name] = {
            "category": get_category(sheet_name),
            "allergens": list(allergen_map.keys()),
            "allergen_sources": {a: list(set(srcs)) for a, srcs in allergen_map.items()},
            "ingredient_count": len(set(raw)),
        }
    return result


# ─────────────────────────────────────────────
# 5. Excel 보고서 생성
# ─────────────────────────────────────────────

C_GREEN = "2D6A4F"
C_GREEN_LIGHT = "D8F3DC"
C_RED = "E63946"
C_RED_LIGHT = "FFE5E5"
C_GRAY = "F8F9FA"
C_BORDER = "DEE2E6"

def _border(color=C_BORDER):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def create_excel_report(menu_allergens: dict, output_path: Path):
    wb = openpyxl.Workbook()

    _build_summary_sheet(wb.active, menu_allergens)
    matrix_ws = wb.create_sheet("알러지 매트릭스")
    _build_matrix_sheet(matrix_ws, menu_allergens)
    detail_ws = wb.create_sheet("메뉴별 상세")
    _build_detail_sheet(detail_ws, menu_allergens)

    wb.save(output_path)
    print(f"Excel 보고서 저장: {output_path}")


def _build_summary_sheet(ws, menu_allergens: dict):
    ws.title = "알러지 요약"

    # 제목
    ws.merge_cells("A1:F1")
    ws["A1"] = "2026 가을·겨울 메뉴 알러지 현황 (식약처 22종 기준)"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = _fill(C_GREEN)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws["A2"] = f"생성일: {date.today().strftime('%Y년 %m월 %d일')}  |  총 {len(menu_allergens)}개 메뉴"
    ws["A2"].font = Font(italic=True, color="666666", size=10)
    ws.merge_cells("A2:F2")

    headers = ["메뉴명", "카테고리", "알러지 수", "알러지 목록", "주요 유발 재료", "비고"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = _fill("40916C")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _border()
    ws.row_dimensions[3].height = 20

    categories = ["코스", "잠실 단품", "단품", "피클/기타", "기타"]
    menus_sorted = sorted(menu_allergens.items(), key=lambda x: (categories.index(x[1]["category"]) if x[1]["category"] in categories else 99, x[0]))

    for r, (name, info) in enumerate(menus_sorted, 4):
        allergen_list = ", ".join(info["allergens"]) if info["allergens"] else "해당 없음"
        sources_flat = []
        for srcs in info["allergen_sources"].values():
            sources_flat.extend(srcs[:1])
        source_str = ", ".join(sources_flat[:4]) + ("..." if len(sources_flat) > 4 else "")

        row_fill = _fill("F8FFF9") if r % 2 == 0 else _fill("FFFFFF")
        values = [name, info["category"], len(info["allergens"]), allergen_list, source_str, ""]

        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.fill = row_fill
            cell.border = _border()
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if c == 3 and v > 0:
                cell.font = Font(bold=True, color=C_RED)
                cell.fill = _fill(C_RED_LIGHT)

        ws.row_dimensions[r].height = 28

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 40
    ws.column_dimensions["F"].width = 12
    ws.freeze_panes = "A4"


def _build_matrix_sheet(ws, menu_allergens: dict):
    ws["A1"] = "메뉴명"
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = _fill(C_GREEN)
    ws["A1"].border = _border()
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 70

    for col, allergen in enumerate(ALLERGEN_NAMES, 2):
        cell = ws.cell(row=1, column=col, value=allergen)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = _fill("40916C")
        cell.border = _border()
        cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=60, wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 5

    categories = ["코스", "잠실 단품", "단품", "피클/기타", "기타"]
    menus_sorted = sorted(menu_allergens.items(), key=lambda x: (categories.index(x[1]["category"]) if x[1]["category"] in categories else 99, x[0]))

    for r, (name, info) in enumerate(menus_sorted, 2):
        name_cell = ws.cell(row=r, column=1, value=name)
        name_cell.border = _border()
        name_cell.alignment = Alignment(vertical="center")
        name_cell.font = Font(size=9)
        ws.row_dimensions[r].height = 18

        for col, allergen in enumerate(ALLERGEN_NAMES, 2):
            cell = ws.cell(row=r, column=col)
            cell.border = _border()
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if allergen in info["allergens"]:
                cell.value = "●"
                cell.fill = _fill(C_RED_LIGHT)
                cell.font = Font(color=C_RED, bold=True, size=10)
            else:
                cell.value = ""
                cell.fill = _fill("FFFFFF") if r % 2 == 0 else _fill(C_GRAY)

    ws.column_dimensions["A"].width = 34
    ws.freeze_panes = "B2"


def _build_detail_sheet(ws, menu_allergens: dict):
    ws["A1"] = "메뉴별 알러지 상세 (유발 재료 포함)"
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = _fill(C_GREEN)
    ws.merge_cells("A1:D1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["메뉴명", "알러지 항목", "유발 재료", "카테고리"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = _fill("40916C")
        cell.border = _border()
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    r = 3
    categories = ["코스", "잠실 단품", "단품", "피클/기타", "기타"]
    menus_sorted = sorted(menu_allergens.items(), key=lambda x: (categories.index(x[1]["category"]) if x[1]["category"] in categories else 99, x[0]))

    for name, info in menus_sorted:
        if not info["allergens"]:
            row_data = [name, "해당 없음", "-", info["category"]]
            for c, v in enumerate(row_data, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = _border()
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.fill = _fill("F0FFF4")
            ws.row_dimensions[r].height = 18
            r += 1
        else:
            for i, allergen in enumerate(info["allergens"]):
                sources = ", ".join(info["allergen_sources"].get(allergen, []))
                row_fill = _fill(C_RED_LIGHT) if i == 0 else _fill("FFF5F5")
                row_data = [name if i == 0 else "", allergen, sources, info["category"] if i == 0 else ""]
                for c, v in enumerate(row_data, 1):
                    cell = ws.cell(row=r, column=c, value=v)
                    cell.border = _border()
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    cell.fill = row_fill
                    if c == 2:
                        cell.font = Font(bold=True, color=C_RED, size=9)
                ws.row_dimensions[r].height = 18
                r += 1

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 12
    ws.freeze_panes = "A3"


# ─────────────────────────────────────────────
# 6. 실행
# ─────────────────────────────────────────────

def main():
    Path("data").mkdir(exist_ok=True)

    print("Excel 파싱 중...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    menu_allergens = build_menu_allergens(wb)

    json_path = Path("data/menu_allergens.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(menu_allergens, f, ensure_ascii=False, indent=2)
    print(f"JSON 저장: {json_path}  ({len(menu_allergens)}개 메뉴)")

    xlsx_path = Path("data/allergen_report.xlsx")
    create_excel_report(menu_allergens, xlsx_path)

    print("\n=== 추출 결과 요약 ===")
    for name, info in menu_allergens.items():
        allergen_str = ", ".join(info["allergens"]) if info["allergens"] else "없음"
        print(f"  [{info['category']}] {name}")
        print(f"    알러지: {allergen_str}")


if __name__ == "__main__":
    main()
